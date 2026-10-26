"""空运代理对账的数据投影、独立结算、原子性及五个接口契约回归。"""
import json
import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.deps import get_current_active_user
from app.api.financial_audit import build_air_financial_audits, get_air_financial_audits
from app.api.reconciliation_agent import (
    _load_actual_flights,
    batch_confirm_agent_settlement, cancel_agent_settlement,
    confirm_agent_settlement, export_agent_reconciliation_list,
    get_agent_reconciliation_list, router,
)
from app.api.reconciliation_airline import (
    cancel_airline_settlement, confirm_airline_settlement, get_airline_reconciliation_list,
)
from app.api.reconciliation_delivery import (
    cancel_delivery_settlement, confirm_delivery_settlement, get_delivery_reconciliation_list,
)
from app.api.reconciliation_pickup import (
    cancel_pickup_settlement, confirm_pickup_settlement, get_pickup_reconciliation_list,
)
from app.database import get_db
from app.models.air_financial_audit_data import AirFinancialAuditData
from app.models.billing_time_container import ShenzhenAirBillingTimeContainer
from app.models.china_southern_air_approval import ChinaSouthernAirApprovalData
from app.models.consignment_note import ConsignmentNote
from app.models.csa_departure_manual_data import CsaDepartureManualData
from app.models.csa_departure_tracking import CsaLalamoveInformation
from app.models.customer import Customer
from app.models.departure_manual_data import ShenzhenAirDepartureManualData
from app.models.peer_air_manual_data import PeerAirDepartureManualData
from app.models.transit_loading import ShenzhenAirBookingExport
from app.models.waybill import Waybill
from app.schemas.financial_audit import AirFinancialAuditQuery, ReceivableResponse
from app.schemas.reconciliation_agent import (
    AgentBatchSettleRequest, AgentReconciliationExportRequest, AgentReconciliationQuery,
)
from app.schemas.reconciliation_airline import AirlineReconciliationQuery
from app.schemas.reconciliation_delivery import DeliveryReconciliationQuery
from app.schemas.reconciliation_pickup import PickupReconciliationQuery


class AgentReconciliationTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
        )
        for model in (
            ShenzhenAirBookingExport, ShenzhenAirDepartureManualData, ShenzhenAirBillingTimeContainer,
            ChinaSouthernAirApprovalData, CsaDepartureManualData, CsaLalamoveInformation,
            ConsignmentNote, PeerAirDepartureManualData, AirFinancialAuditData, Customer, Waybill,
        ):
            model.__table__.create(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.user = SimpleNamespace(id=7, name="测试财务")
        self.db.add(Customer(
            id=7, company_name="测试客户", rate=2, contact_person="张三", contact_phone="13800000000",
            document_fee=10, settlement_cycle=3,
        ))
        self.db.add_all([
            ShenzhenAirBookingExport(
                id=101, prefix="479", waybill_number="62200552", routing="SZX-CGQ",
                flight_date="2026-10-21", quantity="2", chargeable_weight="100",
                freight_rate="2.5", air_freight="250", fuel_surcharge="25",
            ),
            ShenzhenAirDepartureManualData(id=111, booking_export_id=101, customer_name="7", audit_status=2),
            ShenzhenAirBookingExport(id=102, prefix="479", waybill_number="62200553", flight_date="2026-10-21"),
            ShenzhenAirDepartureManualData(id=112, booking_export_id=102, audit_status=1),
            ChinaSouthernAirApprovalData(
                id=201, flight_info="CZ1234/2026-10-22/SZX-HFE", waybill_number="784-12345678",
                chargeable_weight="20", ref_rate="2", ref_freight="40",
            ),
            CsaDepartureManualData(id=211, approval_data_id=201, customer_name="7", audit_status=2),
            ConsignmentNote(
                id=301, transport_type="0", company_name="同行代理", consignment_date=date(2026, 10, 23),
                airline="国航", form_data=json.dumps({"origin_station": "SZX", "quantity": "3"}),
            ),
            PeerAirDepartureManualData(
                id=311, consignment_note_id=301, waybill_number="999-12345678", customer_name="7", audit_status=2,
            ),
            # Hand-entered fee-desk entries identify themselves by fa.id, including legacy source_id=0.
            AirFinancialAuditData(
                id=401, source_type="shenzhen_air", source_id=0, financial_audit_status=2,
                receivable_data={"waybill_number": "479-MANUAL", "flight_date": "2026-10-21",
                                 "airline": "深航", "customer_name": "7", "freight": "30"},
                payable_data={"air_freight": "20"},
            ),
            AirFinancialAuditData(
                id=402, source_type="peer_air", source_id=402, financial_audit_status=2,
                receivable_data={"waybill_number": "999-MANUAL", "flight_date": "2026-10-24"},
            ),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def approve_sources(self, exclude=()):
        """Prepare financially approved rows; business audit values stay untouched."""
        for kind, source_id in (("shenzhen_air", 101), ("china_southern_air", 201), ("peer_air", 301)):
            if (kind, source_id) in exclude:
                continue
            record = self.db.query(AirFinancialAuditData).filter_by(source_type=kind, source_id=source_id).first()
            if record is None:
                record = AirFinancialAuditData(source_type=kind, source_id=source_id)
                self.db.add(record)
            record.financial_audit_status = 2
        self.db.commit()

    async def listing(self, **filters):
        return (await get_agent_reconciliation_list(
            AgentReconciliationQuery(**filters), self.db, self.user,
        )).data

    def settle(self, kind="shenzhen_air", source_id=101):
        return confirm_agent_settlement(kind, source_id, self.db, self.user)

    async def test_all_sources_manual_identity_and_unsettled_without_extension(self):
        self.approve_sources()
        result = await self.listing(settlement_status=0, financial_audit_status=2)
        self.assertEqual(result["total"], 5)
        keys = {(r["source_type"], r["source_id"]) for r in result["items"]}
        self.assertEqual(keys, {("shenzhen_air", "101"), ("china_southern_air", "201"),
                                ("peer_air", "301"), ("shenzhen_air", "401"), ("peer_air", "402")})
        self.assertEqual(self.db.query(AirFinancialAuditData).count(), 5)  # reads never create extensions
        for row in result["items"]:
            self.assertTrue(set(ReceivableResponse.model_fields).issubset(row))
            self.assertEqual(row["agent_settlement_status"], 0)
            for name in ("actual_flight_number", "actual_pieces", "actual_weight"):
                self.assertIsInstance(row[name], str)
        by_id = {row["source_id"]: row for row in result["items"]}
        self.assertEqual(by_id["101"]["waybill_number"], "479-62200552")
        self.assertEqual(by_id["201"]["waybill_number"], "784-12345678")
        self.assertEqual(by_id["301"]["waybill_number"], "999-12345678")
        self.assertEqual(by_id["101"]["actual_customer_name"], "测试客户")

    async def test_actual_transport_values_use_source_flights_and_final_payable(self):
        self.db.get(ShenzhenAirBookingExport, 101).actual_flight = "ZH9637"
        self.db.get(ChinaSouthernAirApprovalData, 201).actual_flight = "CZ5678"
        note = self.db.get(ConsignmentNote, 301)
        note.form_data = json.dumps({"quantity": "99", "actual_flight": "CA3456"})
        self.db.add_all([
            ShenzhenAirBillingTimeContainer(id=601, booking_export_id=101, quantity="2", weight="20.5"),
            ShenzhenAirBillingTimeContainer(id=602, booking_export_id=101, quantity="3", weight="30.25"),
            CsaLalamoveInformation(id=603, approval_data_id=201, pieces="4", weight="40.5"),
            CsaLalamoveInformation(id=604, approval_data_id=201, pieces="1", weight="10"),
            AirFinancialAuditData(id=501, source_type="peer_air", source_id=301, financial_audit_status=2,
                                  payable_data={"gate_pieces": "6", "transit_weight": "60.75"}),
        ])
        self.db.commit()
        self.approve_sources(exclude=(("shenzhen_air", 101),))
        # Keep the Shenzhen financial record available for the override check below.
        self.db.add(AirFinancialAuditData(id=500, source_type="shenzhen_air", source_id=101, financial_audit_status=2))
        self.db.commit()
        rows = {row["source_id"]: row for row in (await self.listing())["items"]}
        for source_id, expected in (
            ("101", ("ZH9637", "5", "50.75")),
            ("201", ("CZ5678", "5", "50.50")),
            ("301", ("CA3456", "6", "60.75")),
        ):
            with self.subTest(source_id=source_id):
                row = rows[source_id]
                self.assertEqual(tuple(row[k] for k in (
                    "actual_flight_number", "actual_pieces", "actual_weight",
                )), expected)

        financial = self.db.get(AirFinancialAuditData, 500)
        financial.payable_data = {"gate_pieces": 0, "transit_weight": 0}
        financial.receivable_data = {"pieces": "999", "chargeable_weight": "999", "flight_number": "ZH9999"}
        self.db.commit()
        row = (await self.listing(waybill_numbers="479-62200552"))["items"][0]
        self.assertEqual((row["actual_flight_number"], row["actual_pieces"], row["actual_weight"]),
                         ("ZH9637", "0", "0"))
        self.assertEqual((row["pieces"], row["chargeable_weight"], row["flight_number"]),
                         ("999", "999", "ZH9999"))

    async def test_actual_flight_fallbacks_and_missing_peer_values(self):
        self.approve_sources()
        self.db.get(ShenzhenAirBookingExport, 101).billing_flight = "ZH1111"
        note = self.db.get(ConsignmentNote, 301)
        note.flight_number = "CA2222"
        manual = self.db.get(AirFinancialAuditData, 401)
        manual.receivable_data = dict(manual.receivable_data, flight_number="MANUAL-FLIGHT")
        manual.payable_data = dict(manual.payable_data, gate_pieces="7", transit_weight="70")
        # A coincident source ID must not overwrite the manual bill's own flight.
        self.db.add(ShenzhenAirBookingExport(id=401, actual_flight="UNRELATED-FLIGHT"))
        self.db.commit()
        rows = {row["source_id"]: row for row in (await self.listing())["items"]}
        self.assertEqual(rows["101"]["actual_flight_number"], "ZH1111")
        self.assertEqual(rows["201"]["actual_flight_number"], "CZ1234")
        self.assertEqual(rows["301"]["actual_flight_number"], "CA2222")
        self.assertEqual((rows["301"]["actual_pieces"], rows["301"]["actual_weight"]), ("", ""))
        self.assertEqual(tuple(rows["401"][k] for k in (
            "actual_flight_number", "actual_pieces", "actual_weight",
        )), ("MANUAL-FLIGHT", "7", "70"))
        self.assertEqual(tuple(rows["402"][k] for k in (
            "actual_flight_number", "actual_pieces", "actual_weight",
        )), ("", "", ""))

    async def test_actual_flights_are_batched_and_handle_invalid_peer_forms(self):
        forms = ["invalid-json", "[]", '{"actual_flight": {}}', '{"actual_flight": true}']
        self.db.add_all([ConsignmentNote(
            id=1000+i, transport_type="0", company_name="批量代理",
            form_data=forms[i] if i < len(forms) else json.dumps({"actual_flight": f"CA{i}"}),
        ) for i in range(501)])
        self.db.commit()
        with patch.object(self.db, "query", wraps=self.db.query) as queries:
            flights = _load_actual_flights(self.db, {("peer_air", 1000+i) for i in range(501)})
        self.assertEqual(queries.call_count, 2)  # 500 IDs per batch, never one query per row.
        self.assertEqual(len(flights), 501)
        for i in range(len(forms)):
            self.assertEqual(flights[("peer_air", str(1000+i))], "")
        self.assertEqual(flights[("peer_air", "1500")], "CA500")

    async def test_final_receivable_matches_fee_desk_overrides_and_totals(self):
        self.db.add(AirFinancialAuditData(
            id=501, source_type="shenzhen_air", source_id=101, financial_audit_status=2,
            payable_data={"air_freight": "100", "fuel_surcharge": "0"},
            receivable_data={"waybill_number": "479-OVERRIDE", "flight_date": "2026-10-25",
                             "freight": "300", "packaging_fee": "5", "remark": "修改后备注"},
        ))
        self.db.commit()
        canonical = await build_air_financial_audits(AirFinancialAuditQuery(), self.db, self.user, paginate=False)
        expected = next(r["receivable"] for r in canonical.data["items"] if r["source_id"] == "101")
        result = await self.listing(waybill_numbers="479-OVERRIDE", flight_date_start="2026-10-25",
                                    flight_date_end="2026-10-25", airline="深航", customer_name="7")
        self.assertEqual(result["total"], 1)
        row = result["items"][0]
        self.assertEqual({k: row[k] for k in ReceivableResponse.model_fields}, expected)
        self.assertEqual(float(row["document_fee"]), 10)
        self.assertEqual(float(row["total_amount"]), 315)
        self.assertEqual(float(row["gross_profit"]), 215)
        self.assertEqual((await self.listing(waybill_numbers="479-62200552"))["total"], 0)

    async def test_filters_apply_before_pagination(self):
        self.approve_sources()
        self.settle()
        result = await self.listing(flight_date_start="2026-10-21", flight_date_end="2026-10-21",
                                   airline="深航", settlement_status=0, pageSize=1)
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["items"][0]["source_id"], "401")
        result = await self.listing(waybill_numbers="479-62200552, 784-12345678", pageSize=1, page=2)
        self.assertEqual(result["total"], 2)
        self.assertEqual(result["items"][0]["source_id"], "101")
        self.assertEqual((await self.listing(airline="全部"))["total"], 5)
        self.assertEqual((await self.listing(financial_audit_status=2))["total"], 5)
        self.assertEqual((await self.listing(financial_audit_status=0))["total"], 0)

    async def test_settle_cancel_idempotent_and_other_categories_untouched(self):
        fa = AirFinancialAuditData(
            id=501, source_type="shenzhen_air", source_id=101,
            airline_settlement_status=1, pickup_settlement_status=1, delivery_settlement_status=1,
            payable_data={"air_freight": "33"}, receivable_data={"remark": "保持原值"}, financial_audit_status=2,
        )
        self.db.add(fa)
        self.db.commit()
        self.settle()
        timestamp = fa.agent_settlement_time
        self.assertEqual(fa.agent_settlement_auditor_id, 7)
        confirm_agent_settlement("shenzhen_air", 101, self.db, SimpleNamespace(id=8, name="另一个人"))
        self.assertEqual(fa.agent_settlement_time, timestamp)
        self.assertEqual(fa.agent_settlement_auditor_name, "测试财务")
        self.assertEqual((await self.listing(settlement_status=1))["total"], 1)
        cancel_agent_settlement("shenzhen_air", 101, self.db, self.user)
        self.assertEqual(fa.agent_settlement_status, 0)
        self.assertIsNone(fa.agent_settlement_auditor_id)
        self.assertIsNone(fa.agent_settlement_auditor_name)
        self.assertIsNone(fa.agent_settlement_time)
        self.assertEqual((fa.airline_settlement_status, fa.pickup_settlement_status, fa.delivery_settlement_status), (1, 1, 1))
        self.assertEqual(fa.financial_audit_status, 2)
        self.assertEqual(fa.payable_data, {"air_freight": "33"})
        self.assertEqual(fa.receivable_data, {"remark": "保持原值"})

    async def test_batch_deduplication_manual_and_linked_sources(self):
        self.approve_sources()
        items = [{"source_type": k, "source_id": str(i)} for k, i in (
            ("shenzhen_air", 101), ("china_southern_air", 201), ("peer_air", 301),
            ("shenzhen_air", 401), ("peer_air", 402), ("shenzhen_air", 101),
        )]
        response = batch_confirm_agent_settlement(AgentBatchSettleRequest(items=items), self.db, self.user)
        self.assertIn("5", response.msg)
        self.assertEqual(self.db.query(AirFinancialAuditData).count(), 5)
        self.assertEqual((await self.listing(settlement_status=1))["total"], 5)
        self.assertEqual(self.db.get(AirFinancialAuditData, 401).source_id, 0)

    async def test_batch_rejects_missing_draft_nonair_and_wrong_manual_alias_without_writes(self):
        for bad_kind, bad_id in (("shenzhen_air", 999), ("shenzhen_air", 102), ("peer_air", 401)):
            with self.subTest(bad_id=bad_id, bad_kind=bad_kind):
                with self.assertRaises(HTTPException) as error:
                    batch_confirm_agent_settlement(AgentBatchSettleRequest(items=[
                        {"source_type": "shenzhen_air", "source_id": "101"},
                        {"source_type": bad_kind, "source_id": str(bad_id)},
                    ]), self.db, self.user)
                self.assertEqual(error.exception.status_code, 404 if bad_id != 102 else 400)
                self.assertEqual(self.db.query(AirFinancialAuditData).count(), 2)
                self.assertEqual((await self.listing(settlement_status=1))["total"], 0)
        note = self.db.get(ConsignmentNote, 301)
        note.transport_type = "1"
        self.db.commit()
        with self.assertRaises(HTTPException):
            self.settle("peer_air", 301)

    async def test_cancel_without_financial_extension_is_noop(self):
        cancel_agent_settlement("shenzhen_air", 101, self.db, self.user)
        self.assertEqual(self.db.query(AirFinancialAuditData).count(), 2)

    async def test_peer_without_business_extension_requires_only_financial_approval(self):
        note = ConsignmentNote(
            id=303, transport_type="0", company_name="同行代理", customer_name="7",
            consignment_date=date(2026, 10, 25), destination="HFE", flight_number="CA1234",
            airline="国航",
            form_data=json.dumps({
                "waybill_number": "999-12340001", "origin_station": "SZX", "quantity": "2",
                "weight": "20", "door_pickup_company": "同行提货", "door_pickup_fee": "15",
                "delivery_company": "同行派送", "delivery_fee": "25", "airport_pickup_fee": "5",
            }),
        )
        self.db.add(note)
        self.db.commit()

        financial = await get_air_financial_audits(
            AirFinancialAuditQuery(airline_type="peer_air", pageSize=200), self.db, self.user,
        )
        peer_item = next(item for item in financial.data["items"] if item["source_id"] == "303")
        self.assertEqual(peer_item["payable"]["billing_pieces"], "2")
        self.assertEqual(peer_item["payable"]["billing_weight"], "20")
        self.assertEqual(peer_item["receivable"]["pieces"], "2")
        self.assertEqual(peer_item["receivable"]["chargeable_weight"], "")
        self.assertEqual((await self.listing(waybill_numbers="999-12340001"))["total"], 0)
        self.db.add(AirFinancialAuditData(source_type="peer_air", source_id=303, financial_audit_status=2))
        self.db.commit()
        listing = await self.listing(waybill_numbers="999-12340001")
        self.assertEqual(listing["total"], 1)
        self.assertEqual(tuple(listing["items"][0][k] for k in (
            "actual_flight_number", "actual_pieces", "actual_weight",
        )), ("CA1234", "", ""))

        for query, list_function, status_field in (
            (AirlineReconciliationQuery, get_airline_reconciliation_list, "airline_settlement_status"),
            (PickupReconciliationQuery, get_pickup_reconciliation_list, "pickup_settlement_status"),
            (DeliveryReconciliationQuery, get_delivery_reconciliation_list, "delivery_settlement_status"),
        ):
            with self.subTest(status_field=status_field):
                for filters in ({}, {"financial_audit_status": 2, "settlement_status": 0}):
                    response = list_function(query(**filters), self.db, self.user)
                    rows = response.data["items"]
                    row = next(item for item in rows if item["source_id"] == "303")
                    self.assertEqual(row["financial_audit_status"], 2)
                    self.assertEqual(row[status_field], 0)

        self.settle("peer_air", 303)
        self.assertEqual((await self.listing(waybill_numbers="999-12340001", settlement_status=1))["total"], 1)
        self.assertIsNone(self.db.query(PeerAirDepartureManualData).filter_by(consignment_note_id=303).first())

    async def test_shenzhen_slash_flight_date_is_included_in_payable_reconciliation_ranges(self):
        self.db.add_all([
            ShenzhenAirBookingExport(
                id=103, prefix="479", waybill_number="62200554", routing="SZX-CGQ",
                flight_date="2026/10/26", quantity="1", weight="10",
            ),
            ShenzhenAirDepartureManualData(
                id=113, booking_export_id=103, customer_name="7", audit_status=0,
                door_pickup_company="测试提货", delivery_company="测试派送", delivery_fee="20",
            ),
            AirFinancialAuditData(source_type="shenzhen_air", source_id=103, financial_audit_status=2),
        ])
        self.db.commit()

        for query, list_function, extra_filters in (
            (AirlineReconciliationQuery, get_airline_reconciliation_list, {}),
            (PickupReconciliationQuery, get_pickup_reconciliation_list, {"pickup_company": "测试提货"}),
            (DeliveryReconciliationQuery, get_delivery_reconciliation_list, {"delivery_company": "测试派送"}),
        ):
            with self.subTest(query=query.__name__):
                response = list_function(query(
                    flight_date_start="2026-10-26", flight_date_end="2026-10-26",
                    pageSize=100, **extra_filters,
                ), self.db, self.user)
                source_ids = {item["source_id"] for item in response.data["items"]}
                self.assertIn("103", source_ids)

    async def test_commit_failure_rolls_back_batch_and_keeps_financial_data(self):
        self.approve_sources()
        with patch.object(self.db, "commit", side_effect=IntegrityError("insert", {}, Exception("conflict"))):
            with self.assertRaises(HTTPException) as error:
                batch_confirm_agent_settlement(AgentBatchSettleRequest(items=[
                    {"source_type": "shenzhen_air", "source_id": "101"},
                    {"source_type": "shenzhen_air", "source_id": "401"},
                ]), self.db, self.user)
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(self.db.query(AirFinancialAuditData).count(), 5)
        self.assertEqual(self.db.get(AirFinancialAuditData, 401).agent_settlement_status, 0)
        self.assertEqual(self.db.get(AirFinancialAuditData, 401).receivable_data["freight"], "30")

    async def test_linked_financial_record_id_cannot_be_used_as_source_id(self):
        self.db.add(AirFinancialAuditData(id=501, source_type="shenzhen_air", source_id=101))
        self.db.commit()
        with self.assertRaises(HTTPException) as error:
            self.settle(source_id=501)
        self.assertEqual(error.exception.status_code, 404)
        self.assertEqual(self.db.get(AirFinancialAuditData, 501).agent_settlement_status, 0)

    async def test_existing_three_reconciliation_operations_preserve_agent_settlement(self):
        self.approve_sources()
        self.settle()
        fa = self.db.query(AirFinancialAuditData).filter_by(source_type="shenzhen_air", source_id=101).one()
        timestamp = fa.agent_settlement_time
        for confirm, cancel in (
            (confirm_airline_settlement, cancel_airline_settlement),
            (confirm_pickup_settlement, cancel_pickup_settlement),
            (confirm_delivery_settlement, cancel_delivery_settlement),
        ):
            confirm("shenzhen_air", "101", self.db, self.user)
            cancel("shenzhen_air", "101", self.db, self.user)
            self.assertEqual(fa.agent_settlement_status, 1)
            self.assertEqual(fa.agent_settlement_time, timestamp)
            self.assertEqual(fa.agent_settlement_auditor_name, "测试财务")

    async def export_rows(self, **values):
        response = await export_agent_reconciliation_list(
            AgentReconciliationExportRequest(**values), self.db, self.user,
        )
        body = b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(BytesIO(body), read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        return rows

    async def test_export_is_unpaged_and_has_every_receivable_field(self):
        self.approve_sources()
        rows = await self.export_rows(pageSize=1, page=2)
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[0], tuple(["序号", "结算状态", "客户公司名称"] + [
            f.description for f in ReceivableResponse.model_fields.values()
        ] + ["实走航班号", "实走件数", "实走重量"]))
        self.assertTrue(all(row[1] == "未结算" for row in rows[1:]))

    async def test_export_actual_columns_match_list_and_preserve_zero(self):
        self.db.get(ShenzhenAirBookingExport, 101).actual_flight = "=SOURCE-FLIGHT"
        self.db.add(AirFinancialAuditData(
            id=501, source_type="shenzhen_air", source_id=101, financial_audit_status=2,
            payable_data={"gate_pieces": "0", "transit_weight": "0"},
        ))
        self.db.commit()
        item = (await self.listing(waybill_numbers="479-62200552"))["items"][0]
        rows = await self.export_rows(selected_items=[{"source_type": "shenzhen_air", "source_id": "101"}])
        for name, label in (("actual_flight_number", "实走航班号"),
                            ("actual_pieces", "实走件数"), ("actual_weight", "实走重量")):
            self.assertEqual(rows[1][rows[0].index(label)], item[name])
        response = await export_agent_reconciliation_list(
            AgentReconciliationExportRequest(waybill_numbers="479-62200552"), self.db, self.user,
        )
        body = b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=False)
        self.assertTrue(all(cell.data_type != "f" for row in workbook.active.rows for cell in row))
        workbook.close()

    async def test_export_selection_filters_and_formula_protection(self):
        fa = self.db.get(AirFinancialAuditData, 401)
        fa.receivable_data = dict(fa.receivable_data, remark="=HYPERLINK(\"https://example.com\")")
        self.db.commit()
        rows = await self.export_rows(airline="深航", selected_items=[
            {"source_type": "shenzhen_air", "source_id": "401"},
            {"source_type": "china_southern_air", "source_id": "201"},
        ])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][rows[0].index("备注(人工填写)")], fa.receivable_data["remark"])
        response = await export_agent_reconciliation_list(
            AgentReconciliationExportRequest(waybill_numbers="479-MANUAL"), self.db, self.user,
        )
        body = b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=False)
        self.assertTrue(all(cell.data_type != "f" for row in workbook.active.rows for cell in row))
        workbook.close()
        with self.assertRaises(HTTPException) as error:
            await self.export_rows(waybill_numbers="不存在")
        self.assertEqual(error.exception.status_code, 400)

    async def test_projection_not_limited_to_fee_desk_page_size(self):
        self.approve_sources()
        self.db.add_all([AirFinancialAuditData(
            id=1000+i, source_type="peer_air", source_id=1000+i, financial_audit_status=2,
            receivable_data={"waybill_number": f"999-{i:08}", "flight_date": "2026-10-21"},
        ) for i in range(205)])
        self.db.commit()
        result = await self.listing(pageSize=1)
        self.assertEqual(result["total"], 210)
        self.assertEqual(len(await self.export_rows(pageSize=1)), 211)

    async def test_http_routes_search_contract_validation_and_auth(self):
        self.approve_sources()
        app = FastAPI()
        prefix = "/api/v1/reconciliation/air-agent"
        app.include_router(router, prefix=prefix)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_active_user] = lambda: self.user
        with TestClient(app) as client:
            schema = client.get("/openapi.json").json()
            self.assertEqual(sum(len(p) for p in schema["paths"].values()), 5)
            params = {p["name"] for p in schema["paths"][prefix+"/air"]["get"]["parameters"]}
            self.assertEqual(params, set(AirlineReconciliationQuery.model_fields))
            response = client.get(prefix+"/air", params={"waybill_numbers": "479-62200552"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["data"]["total"], 1)
            for name in ("actual_flight_number", "actual_pieces", "actual_weight"):
                self.assertIn(name, schema["components"]["schemas"]["AgentReconciliationItemResponse"]["properties"])
                self.assertIsInstance(response.json()["data"]["items"][0][name], str)
            for params in ({"flight_date_start": "bad"}, {"settlement_status": 3}, {"page": 0}):
                self.assertEqual(client.get(prefix+"/air", params=params).status_code, 422)
            self.assertEqual(client.get(prefix+"/air", params={
                "flight_date_start": "2026-10-22", "flight_date_end": "2026-10-21",
            }).status_code, 400)
            for path in ("bad/101", "shenzhen_air/0", "shenzhen_air/9223372036854775808"):
                self.assertEqual(client.post(prefix+"/air/"+path+"/settle").status_code, 422)
            self.assertEqual(client.post(prefix+"/air/batch-settle", json={"items": [
                {"source_type": "peer_air", "source_id": "0"},
            ]}).status_code, 422)
            self.assertEqual(client.post(prefix+"/air/shenzhen_air/101/settle").status_code, 200)
            self.assertEqual(client.post(prefix+"/air/shenzhen_air/101/cancel-settlement").status_code, 200)
            self.assertEqual(client.post(prefix+"/air/batch-settle", json={"items": []}).status_code, 200)
            response = client.post(prefix+"/air/export", json={"airline": "深航"})
            self.assertEqual(response.status_code, 200)
            self.assertIn("spreadsheetml", response.headers["content-type"])
            del app.dependency_overrides[get_current_active_user]
            for method, path, body in (
                ("GET", "/air", None), ("POST", "/air/shenzhen_air/101/settle", None),
                ("POST", "/air/shenzhen_air/101/cancel-settlement", None),
                ("POST", "/air/batch-settle", {"items": []}), ("POST", "/air/export", {}),
            ):
                self.assertIn(client.request(method, prefix+path, json=body).status_code, (401, 403))


if __name__ == "__main__":
    unittest.main()
