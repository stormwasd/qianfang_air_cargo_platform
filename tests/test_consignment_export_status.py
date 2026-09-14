import io
import unittest
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import MagicMock

from openpyxl import load_workbook

from app.api.cost_service import export_cost_consignments_to_excel
from app.api.customer_service import export_consignments_to_excel
from app.models.cost_service import CostConsignment
from app.models.customer_service import ConsignmentInfo
from app.schemas.cost_service import CostExportExcelRequest
from app.schemas.customer_service import ExportExcelRequest


class ConsignmentExportStatusTests(unittest.IsolatedAsyncioTestCase):
    async def export(self, endpoint, request_type, records):
        db = MagicMock()
        db.query.return_value.filter.return_value.order_by.return_value.all.return_value = records
        response = await endpoint(
            request_type(ids=["1", "2"]), current_user=None, db=db,
        )
        workbook = load_workbook(io.BytesIO(response.body), data_only=True)
        self.addCleanup(workbook.close)
        return workbook.active

    async def test_both_exports_use_their_own_status_and_preserve_consignment_columns(self):
        fields = {
            "id": 1,
            "create_time": datetime(2026, 9, 14, 10, 20, 30),
            "internal_doc_id": "DOC-1",
            "warehouse_entry_date": date(2026, 9, 13),
            "customer_name": "customer",
            "origin_destination": "SZX-NRT",
            "customs_declaration": "customs",
            "bill_of_lading": "2-0",
            "flight_date": date(2026, 9, 16),
            "flight_no": "MF809",
            "flight_doc_no": "731-90064074",
            "pieces": 36,
            "actual_weight": Decimal("123.25"),
            "chargeable_weight": Decimal("125.00"),
            "volume": Decimal("1.25"),
            "first_leg_weight": Decimal("120.50"),
            "agent": "agent",
            "remark": "remark",
        }
        expected_headers = [
            "状态", "制单时间", "内部单据ID", "进仓日期", "客户名称",
            "始发站-目的站", "报关", "提单", "航班日期", "航班号",
            "航班单号", "件数", "实际重量(kg)", "计费重量(kg)",
            "体积(m³)", "一程重量(kg)", "代理", "备注",
        ]
        expected_values = [
            "2026-09-14 10:20:30", "DOC-1", "2026-09-13", "customer",
            "SZX-NRT", "customs", "直单", "2026-09-16", "MF809",
            "731-90064074", 36, 123.25, 125, 1.25, 120.5, "agent", "remark",
        ]
        cases = (
            (export_consignments_to_excel, ExportExcelRequest, ConsignmentInfo, 2, 18, 1),
            (export_cost_consignments_to_excel, CostExportExcelRequest, CostConsignment, 4, 117, 0),
        )
        for endpoint, request_type, model, data_row, column_count, status in cases:
            with self.subTest(endpoint=endpoint.__name__):
                record = model(**fields, status=status)
                second = model(id=2, status=1 - status)
                sheet = await self.export(endpoint, request_type, [record, second])
                self.assertEqual(sheet.max_column, column_count)
                self.assertEqual(sheet["A1"].value, "状态")
                self.assertEqual(sheet.cell(data_row, 1).value, "已提交" if status == 1 else "未提交")
                self.assertEqual(sheet.cell(data_row + 1, 1).value, "未提交" if status == 1 else "已提交")
                self.assertEqual(
                    [sheet.cell(data_row, c).value for c in range(2, 19)],
                    expected_values,
                )
                self.assertEqual(record.status, status)
                if data_row == 2:
                    self.assertEqual([cell.value for cell in sheet[1]], expected_headers)
                    for column in (2, 4, 9, 12, 13, 14, 15, 16):
                        self.assertEqual(sheet.cell(data_row, column).alignment.horizontal, "center")
                    self.assertEqual(sheet.cell(data_row, 3).alignment.horizontal, "left")
                else:
                    self.assertIn("A1:A3", {str(r) for r in sheet.merged_cells.ranges})
                    self.assertEqual(
                        [sheet.cell(3, c).value for c in range(2, 19)], expected_headers[1:],
                    )

    async def test_empty_result_still_exports_status_headers(self):
        for endpoint, request_type, column_count, header_rows in (
            (export_consignments_to_excel, ExportExcelRequest, 18, 1),
            (export_cost_consignments_to_excel, CostExportExcelRequest, 117, 3),
        ):
            with self.subTest(endpoint=endpoint.__name__):
                sheet = await self.export(endpoint, request_type, [])
                self.assertEqual(sheet["A1"].value, "状态")
                self.assertEqual(sheet.max_column, column_count)
                self.assertEqual(sheet.max_row, header_rows)


if __name__ == "__main__":
    unittest.main()
