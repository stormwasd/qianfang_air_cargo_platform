"""隔离数据库验证客户新增录入、完整流转、提交锁定及下载内容。"""
import io
import unittest
from copy import deepcopy
from decimal import Decimal
from types import SimpleNamespace

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.api import customer_service, cost_service
from app.api.deps import get_current_active_user
from app.database import get_db
from app.models.customer_service import ConsignmentInfo, ConsignmentRegistration
from app.models.cost_service import CostConsignment, CostRegistration
from app.models.consignment_operation_log import ConsignmentOperationLog
from app.services.consignment_financial import FINANCIAL_FIELD_MAP
from app.main import register_exception_handlers


FINANCIAL = {
    "receivables": {
        "buy_order_fee": 300, "document_fee": 150, "customs_fee": 750,
        "fuel_rate": 0.65, "fuel_weight_method": "1", "receivable_fuel_fee": 65,
        "tc_rate": 0.3, "tc_express_fee": 36, "front_warehouse_rate": 0.2,
        "warehouse_ground_fee": 24, "remark": "应收单独备注",
    },
    "payables": {
        "intl_air": {
            "outsource_unit": "国际代理", "weight": 105, "flight_doc_no": "999-00000001",
            "buy_order_fee": 280, "document_fee": 140, "customs_fee": 700,
            "fuel_rate": 0.5, "fuel_weight_method": "2", "borrow_magnetic_fuel_pickup_fee": 60,
            "tc_rate": 0.25, "tc_network_disposal_fee": 30,
            "front_warehouse_rate": 0.1, "front_warehouse": 12,
            "flight_date": "2026-10-30", "remark": "国际备注",
        },
        "customs": {"agent": "报关代理", "buy_order_fee": 260, "document_fee": 130, "customs_fee": 650},
        "ground": {"rate": 0.12, "freight": 14.4, "tc_rate": 0.15, "tc_fee": 18, "pickup_fee": 0},
        "trucking": {"flight_doc_no": "ROAD-001", "origin_city": "广州", "destination_city": "深圳", "weight": 120},
        "dom_air": {"volume": 1.234},
        "pay_total": 1784.4,
    },
}


class ConsignmentFinancialFieldsTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        for model in (ConsignmentInfo, ConsignmentRegistration, CostConsignment, CostRegistration, ConsignmentOperationLog):
            model.__table__.create(self.engine)
        self.db = Session(self.engine, autoflush=False)
        app = FastAPI()
        register_exception_handlers(app)
        app.include_router(customer_service.router, prefix="/customer")
        app.include_router(cost_service.router, prefix="/cost")

        async def get_test_db():
            return self.db

        async def get_test_user():
            return SimpleNamespace(id=99, name="操作员")

        app.dependency_overrides[get_db] = get_test_db
        app.dependency_overrides[get_current_active_user] = get_test_user
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        self.db.close()
        self.engine.dispose()

    async def request(self, method, path, payload=None):
        response = await self.client.request(method, path, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["data"]

    async def test_customer_draft_partial_clear_submit_lock_and_cost_reverse_sync(self):
        data = await self.request("POST", "/customer/consignments/draft", {
            "customer_name": "客户", "actual_weight": 100, "customs_declaration": "1买单+3单证", **FINANCIAL,
        })
        record_id = int(data["id"])
        self.assertIsNone(self.db.get(CostConsignment, record_id))
        path = "/customer/consignments/" + data["id"]
        data = await self.request("PUT", path + "/draft", {
            "receivables": {"fuel_rate": 0, "remark": None}, "payables": {"ground": {"tc_fee": None}},
        })
        self.assertEqual(data["receivables"]["buy_order_fee"], 300)
        self.assertEqual(data["receivables"]["fuel_rate"], 0)
        self.assertIsNone(data["receivables"]["remark"])
        data = await self.request("PUT", path, {})
        self.assertEqual(data["status"], 1)
        cost = self.db.get(CostConsignment, record_id)
        self.assertEqual(cost.status, 0)
        self.assertEqual(cost.actual_weight, Decimal("100.00"))
        self.assertEqual(cost.pay_intl_air_weight, Decimal("105.00"))
        self.assertEqual(cost.pay_intl_air_flight_doc_no, "999-00000001")
        self.assertEqual(cost.pay_dom_air_volume, Decimal("1.234"))
        self.assertEqual(cost.receivable_customs_fee, Decimal("750.00"))
        self.assertEqual(cost.receivable_fuel_fee, Decimal("65.00"))  # 不随后端费率 0 回算。
        before = self.db.query(ConsignmentOperationLog).count()
        for suffix in ("", "/draft"):
            response = await self.client.put(path + suffix, json={"customer_name": "非法覆盖", "status": 0})
            self.assertEqual(response.status_code, 409, response.text)
            self.assertEqual(response.json(), {
                "code": 409, "data": None,
                "msg": "已提交的委托信息不能修改或暂存，请在费用登记台维护",
            })
        self.assertEqual(self.db.query(ConsignmentOperationLog).count(), before)
        self.assertEqual(self.db.get(ConsignmentInfo, record_id).customer_name, "客户")
        await self.request("PUT", "/cost/consignments/" + str(record_id), {
            "receivables": {"fuel_rate": 1.23}, "payables": {"trucking": {"destination_city": "北京"}},
        })
        detail = await self.request("GET", path)
        self.assertEqual(detail["receivables"]["fuel_rate"], 1.23)
        self.assertEqual(detail["payables"]["trucking"]["destination_city"], "北京")
        self.assertEqual(detail["status"], 1)

    async def test_templates_and_cost_create_draft_update_list_void_return_new_fields(self):
        for base, payload in (("/customer", FINANCIAL), ("/cost", FINANCIAL)):
            registration = "/consignment-registration" if base == "/customer" else "/cost-registration"
            await self.request("PUT", base + registration, payload)
            result = await self.request("GET", base + registration)
            self.assertEqual(result["receivables"]["fuel_rate"], 0.65)
            self.assertEqual(result["payables"]["trucking"]["origin_city"], "广州")
        for suffix in ("", "/draft"):
            result = await self.request("POST", "/cost/consignments" + suffix, FINANCIAL)
            record_id = result["id"]
            path = "/cost/consignments/" + record_id
            self.assertEqual(result["payables"]["dom_air"]["volume"], 1.234)
            if not suffix:
                customer = await self.request("GET", "/customer/consignments/" + record_id)
                self.assertEqual(customer["payables"]["customs"]["buy_order_fee"], 260)
            result = await self.request("PUT", path + "/draft", {"receivables": {"fuel_rate": None}})
            self.assertIsNone(result["receivables"]["fuel_rate"])
            self.assertEqual(result["receivables"]["buy_order_fee"], 300)
            result = await self.request("PUT", path, {"receivables": {"fuel_rate": 0}})
            self.assertEqual(result["receivables"]["fuel_rate"], 0)
            result = await self.request("GET", "/cost/consignments")
            self.assertTrue(any(item["id"] == record_id and item["receivables"]["fuel_rate"] == 0 for item in result["items"]))
            result = await self.request("PUT", path + "/void")
            self.assertEqual(result["status"], 2)
            self.assertEqual(result["payables"]["ground"]["tc_rate"], 0.15)

    async def test_new_fields_export_with_numeric_zero_blank_and_weight_labels(self):
        payload = deepcopy(FINANCIAL)
        payload["receivables"]["fuel_rate"] = 0
        result = await self.request("POST", "/customer/consignments", payload)
        for base, header_row, data_row, count in (("/customer", 1, 2, 132), ("/cost", 3, 4, 137)):
            response = await self.client.post(base + "/consignments/export-excel", json={"ids": [result["id"]]})
            self.assertEqual(response.status_code, 200, response.text)
            workbook = load_workbook(io.BytesIO(response.content), data_only=True)
            sheet = workbook.active
            self.assertEqual(sheet.max_column, count)
            offset = 17 if base == "/customer" else 93
            titles = [sheet.cell(header_row, i).value for i in range(offset + 1, count + 1)]
            values = [sheet.cell(data_row, i).value for i in range(offset + 1, count + 1)]
            self.assertIn("广州", values)
            self.assertIn("ROAD-001", values)
            self.assertIn(1.234, values)
            self.assertIn("实际重量", values)
            self.assertIn("计费重量", values)
            fuel_column = next(i for i, title in enumerate(titles) if title.endswith("燃油费率"))
            self.assertEqual(values[fuel_column], 0)
            self.assertEqual(sheet.cell(data_row, offset + fuel_column + 1).data_type, "n")
            workbook.close()

    async def test_invalid_weight_method_rejected_and_legacy_payload_remains_supported(self):
        response = await self.client.post("/customer/consignments", json={"receivables": {"fuel_weight_method": "3"}})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.db.query(ConsignmentInfo).count(), 0)
        data = await self.request("POST", "/customer/consignments", {"pieces": 8, "actual_weight": 10})
        self.assertIsNone(data["receivables"]["fuel_rate"])
        self.assertEqual(self.db.get(CostConsignment, int(data["id"])).pay_intl_air_pieces, 8)

    def test_every_public_financial_field_maps_to_a_real_column(self):
        for path, fields in FINANCIAL_FIELD_MAP.items():
            for name, column in fields.items():
                self.assertIn(column, CostConsignment.__table__.columns, (path, name))
