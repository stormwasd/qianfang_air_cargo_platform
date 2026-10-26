"""汽运代理对账：数据来源、日期筛选、独立结算、导出与HTTP契约。"""
import json
import unittest
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.consignment_notes import (
    audit_consignment_note, financial_audit_consignment_note, get_consignment_notes,
)
from app.api.deps import get_current_active_user
from app.api.reconciliation_road_agent import (
    batch_confirm_road_agent_settlement, cancel_road_agent_settlement,
    confirm_road_agent_settlement, export_road_agent_reconciliation_list,
    get_road_agent_reconciliation_list, router,
)
from app.database import get_db
from app.models.consignment_note import ConsignmentNote
from app.models.customer import Customer
from app.models.peer_road_manual_data import PeerRoadDepartureManualData
from app.schemas.consignment_note import (
    ConsignmentNoteFinancialAudit, ConsignmentNoteQuery, PeerAirDepartureManualDataUpsert,
)
from app.schemas.reconciliation_road_agent import (
    RoadAgentBatchSettleRequest, RoadAgentReceivableResponse,
    RoadAgentReconciliationExportRequest, RoadAgentReconciliationQuery,
)


class RoadAgentTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        for model in (ConsignmentNote, PeerRoadDepartureManualData, Customer):
            model.__table__.create(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.user = SimpleNamespace(id=7, name="测试财务")
        self.db.add(Customer(id=7, company_name="测试客户", rate=2, contact_person="张三", contact_phone="13800000000"))
        self.form = {
            "transport_date": "2026-10-21", "cargo_name": "五金件", "origin_city": "SZX",
            "origin_address": "始发详细地址", "destination_city": "CGQ", "destination_address": "目的详细地址",
            "consignee_info": "收货人", "contact_phone": "0013800000000", "quantity": 2,
            "weight": 100.5, "chargeable_weight": "120.00", "volume": "0.8", "vehicle_type": "厢式货车", "rate": 2.5,
            "total_freight": "250.00", "other_fees": "25.00", "total_amount": "300.00",
        }
        self.add_note(101, customer="7", form=self.form)
        self.add_note(102, customer="普通客户", destination="SZX", day=date(2026, 10, 20), approved=False)
        self.add_note(103, customer="百分%客户", destination="SZ_", day=date(2026, 10, 19))
        self.add_note(104, day=date(2026, 9, 20))
        self.add_note(105, day=date(2026, 10, 22))
        self.add_note(106, transport_type="0")
        self.db.add(PeerRoadDepartureManualData(
            id=202, consignment_note_id=102, audit_status=1, financial_audit_status=2,
            financial_auditor_name="原审核人", agent_settlement_status=0,
        ))
        self.db.commit()

    def add_note(self, source_id, *, customer="原客户", destination="CGQ", day=date(2026, 10, 21), transport_type="1", form=None, approved=True):
        self.db.add(ConsignmentNote(
            id=source_id, transport_type=transport_type, customer_name=customer,
            consignment_date=day, destination=destination,
            form_data=json.dumps(form or {}, ensure_ascii=False),
        ))
        if approved and transport_type == "1":
            self.db.add(PeerRoadDepartureManualData(
                id=10000 + source_id, consignment_note_id=source_id, financial_audit_status=2,
            ))

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def listing(self, **kwargs):
        query = RoadAgentReconciliationQuery(date_start=date(2026, 9, 21), date_end=date(2026, 10, 21), **kwargs)
        return get_road_agent_reconciliation_list(query, self.db, self.user).data

    def settle(self, source_id=101, user=None):
        return confirm_road_agent_settlement(source_id, self.db, user or self.user)

    def record(self, source_id):
        return self.db.query(PeerRoadDepartureManualData).filter_by(consignment_note_id=source_id).first()

    async def test_approved_fields_from_fee_desk_without_recalculation_or_business_gate(self):
        result = self.listing()
        self.assertEqual(result["total"], 3)
        row = result["items"][0]
        self.assertEqual(row["source_id"], "101")
        self.assertEqual(row["actual_customer_name"], "测试客户")
        for field in RoadAgentReceivableResponse.model_fields:
            expected = "7" if field == "customer_name" else str(self.form[field])
            self.assertEqual(row[field], expected, field)
        # This amount is intentionally different from freight + other fees.
        self.assertEqual(row["total_amount"], "300.00")
        self.assertEqual(row["agent_settlement_status"], 0)
        self.assertEqual(self.record(101).audit_status, 0)
        self.assertEqual(self.listing(page=2, pageSize=1)["items"][0]["source_id"], "102")

    async def test_defaults_use_beijing_date_and_30_day_window(self):
        with patch("app.schemas.reconciliation_road_agent.get_china_now", return_value=datetime(2026, 10, 21, 0, 1)):
            query = RoadAgentReconciliationQuery()
        self.assertEqual(query.date_start, date(2026, 9, 21))
        self.assertEqual(query.date_end, date(2026, 10, 21))

    async def test_search_by_customer_id_name_city_and_literal_wildcards(self):
        for term in ("7", "测试", "测试客户"):
            self.assertEqual(self.listing(customer_name=term)["items"][0]["source_id"], "101")
        self.assertEqual(self.listing(customer_name="普通")["total"], 1)
        self.assertEqual(self.listing(customer_name="%")["total"], 1)
        self.assertEqual(self.listing(destination_city="_")["items"][0]["source_id"], "103")
        self.assertEqual(self.listing(destination_city="CG")["total"], 1)
        self.assertEqual(self.listing(customer_name="无匹配")["total"], 0)

    async def test_settlement_filter_handles_missing_null_and_pagination(self):
        record = self.record(102)
        record.agent_settlement_status = None
        record.receivable_agent_settlement_status = 1
        self.db.commit()
        self.settle(101)

        self.assertEqual(self.listing()["total"], 3)
        settled = self.listing(settlement_status=1)
        self.assertEqual(settled["total"], 1)
        self.assertEqual(settled["items"][0]["source_id"], "101")
        for page, source_id in ((1, "102"), (2, "103")):
            result = self.listing(settlement_status=0, page=page, pageSize=1)
            self.assertEqual(result["total"], 2)
            self.assertEqual(result["items"][0]["source_id"], source_id)
            self.assertEqual(result["items"][0]["agent_settlement_status"], 0)
        self.assertEqual(self.listing(settlement_status=0, customer_name="7")["total"], 0)
        self.assertEqual(self.listing(settlement_status=1, destination_city="CG")["total"], 1)
        self.assertEqual(self.db.query(PeerRoadDepartureManualData).count(), 5)
        self.assertEqual(self.record(103).agent_settlement_status, 0)

        cancel_road_agent_settlement(101, self.db, self.user)
        self.assertEqual(self.listing(settlement_status=1)["total"], 0)
        self.assertEqual(self.listing(settlement_status=0)["total"], 3)

    async def test_missing_total_is_not_fabricated_and_numbers_preserve_zero(self):
        note = self.db.get(ConsignmentNote, 101)
        note.form_data = json.dumps({"quantity": 0, "total_freight": 0, "other_fees": "25"})
        self.db.commit()
        row = self.listing()["items"][0]
        self.assertIsNone(row["total_amount"])
        self.assertEqual(row["quantity"], "0")
        self.assertEqual(row["total_freight"], "0")
        self.assertEqual(row["contact_phone"], None)

    async def test_customer_ids_match_exactly_without_numeric_coercion_or_precision_loss(self):
        customer_id = 371229268321112064
        self.db.add(Customer(id=customer_id, company_name="大ID客户", rate=2,
                             contact_person="张三", contact_phone="13800000000"))
        for source_id, customer in ((107, str(customer_id)), (108, str(customer_id + 1)),
                                    (109, "7abc"), (110, "007"), (111, "7.0")):
            self.add_note(source_id, customer=customer)
        self.db.commit()
        rows = {row["source_id"]: row for row in self.listing()["items"]}
        self.assertEqual(rows["107"]["actual_customer_name"], "大ID客户")
        self.assertEqual(rows["108"]["actual_customer_name"], str(customer_id + 1))
        for source_id, customer in ((109, "7abc"), (110, "007"), (111, "7.0")):
            self.assertEqual(rows[str(source_id)]["actual_customer_name"], customer)
        self.assertEqual(self.listing(customer_name="大ID")["total"], 1)

    async def test_chargeable_weight_preserves_saved_value_without_using_gross_weight(self):
        note = self.db.get(ConsignmentNote, 101)
        for value, expected in ((0, "0"), (120.5, "120.5"), ("120.00", "120.00"),
                                (None, None), (True, None), ({"value": 120}, None)):
            with self.subTest(value=value):
                note.form_data = json.dumps({**self.form, "chargeable_weight": value})
                self.db.commit()
                self.assertEqual(self.listing()["items"][0]["chargeable_weight"], expected)
        note.form_data = json.dumps({"weight": 100})
        self.db.commit()
        row = self.listing()["items"][0]
        self.assertEqual(row["weight"], "100")
        self.assertIsNone(row["chargeable_weight"])

    async def test_actual_fee_desk_sample_matches_before_and_after_audits(self):
        source_id = 371229268321112064
        form = {
            "flight_date": "2026-10-21", "transport_date": "2026-10-21",
            "quantity": 1, "weight": 1, "volume": 1, "cargo_name": "1",
            "other_fees": 1, "vehicle_type": "1", "total_freight": 1,
            "origin_city": "北京市-市辖区-西城区", "origin_address": "11 Delamere Street",
            "destination_city": "河北省-保定市-徐水区", "destination_address": "11 Delamere Street",
            "consignee_info": "1", "contact_phone": 1380000000,
        }
        self.add_note(source_id, customer="Tester", destination=form["destination_city"], form=form, approved=False)
        self.db.commit()
        fee_desk = await get_consignment_notes(ConsignmentNoteQuery(
            transport_type="1", is_financial=True, customer_name="Tester",
            date_start="2026-10-21", date_end="2026-10-22",
        ), self.user, self.db)
        source = fee_desk.data["items"][0]
        self.assertEqual(source["form_data"], form)
        self.assertIsNone(source["manual_data"])
        self.assertEqual(self.listing(customer_name="Tester")["total"], 0)
        for action in ("save", "submit"):
            await financial_audit_consignment_note(ConsignmentNoteFinancialAudit(
                consignment_note_id=str(source_id),
            ), action, self.user, self.db)
            if action == "save":
                self.assertEqual(self.listing(customer_name="Tester")["total"], 0)
                continue
            row = self.listing(customer_name="Tester")["items"][0]
            self.assertEqual(row["source_id"], str(source_id))
            self.assertIsNone(row["rate"])
            self.assertIsNone(row["total_amount"])
            for field, value in form.items():
                if field != "flight_date":
                    self.assertEqual(row[field], str(value), field)
        self.assertEqual(json.loads(self.db.get(ConsignmentNote, source_id).form_data), form)

    async def test_idempotent_settlement_and_cancellation_preserve_original_data(self):
        self.settle()
        record = self.record(101)
        timestamp = record.agent_settlement_time
        self.settle(user=SimpleNamespace(id=8, name="其他财务"))
        self.assertEqual(record.agent_settlement_auditor_name, "测试财务")
        self.assertEqual(record.agent_settlement_time, timestamp)
        self.assertEqual(record.audit_status, 0)
        self.assertEqual(record.financial_audit_status, 2)
        self.assertEqual(json.loads(self.db.get(ConsignmentNote, 101).form_data), self.form)
        cancel_road_agent_settlement(101, self.db, self.user)
        self.assertEqual(record.agent_settlement_status, 0)
        self.assertIsNone(record.agent_settlement_auditor_id)
        self.assertIsNone(record.agent_settlement_time)
        cancel_road_agent_settlement(103, self.db, self.user)
        self.assertEqual(self.record(103).agent_settlement_status, 0)

    async def test_batch_deduplicates_and_rejects_invalid_atomically(self):
        req = RoadAgentBatchSettleRequest(items=[{"source_id": "101"}, {"source_id": "102"}, {"source_id": "101"}])
        result = batch_confirm_road_agent_settlement(req, self.db, self.user)
        self.assertIn("2 条", result.msg)
        self.assertEqual(self.record(102).financial_audit_status, 2)
        self.assertEqual(self.record(102).audit_status, 1)
        for invalid in (106, 999):
            with self.assertRaises(HTTPException) as error:
                batch_confirm_road_agent_settlement(RoadAgentBatchSettleRequest(items=[
                    {"source_id": "103"}, {"source_id": str(invalid)},
                ]), self.db, self.user)
            self.assertEqual(error.exception.status_code, 404)
            self.assertEqual(self.record(103).agent_settlement_status, 0)

    async def test_commit_failure_rolls_back(self):
        with patch.object(self.db, "commit", side_effect=IntegrityError("insert", {}, Exception("duplicate"))):
            with self.assertRaises(HTTPException) as error:
                self.settle()
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(self.record(101).agent_settlement_status, 0)

    async def test_existing_document_and_financial_audits_preserve_settlement(self):
        self.settle()
        await audit_consignment_note(PeerAirDepartureManualDataUpsert(consignment_note_id="101"), "submit", self.user, self.db)
        await financial_audit_consignment_note(ConsignmentNoteFinancialAudit(consignment_note_id="101"), "submit", self.user, self.db)
        record = self.record(101)
        self.assertEqual(record.agent_settlement_status, 1)
        self.assertEqual(record.audit_status, 2)
        self.assertEqual(record.financial_audit_status, 2)

    async def export(self, **kwargs):
        response = export_road_agent_reconciliation_list(RoadAgentReconciliationExportRequest(
            date_start="2026-09-21", date_end="2026-10-21", **kwargs,
        ), self.db, self.user)
        data = b"".join([chunk async for chunk in response.body_iterator])
        return load_workbook(BytesIO(data)).active

    async def test_export_all_selected_intersection_formula_and_phone(self):
        self.settle()
        note = self.db.get(ConsignmentNote, 101)
        form = dict(self.form, cargo_name="=1+1")
        note.form_data = json.dumps(form)
        self.db.commit()
        sheet = await self.export(pageSize=1)
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.cell(2, 2).value, "已结算")
        self.assertEqual(sheet.cell(2, 4).value, "测试客户")
        self.assertEqual(sheet.cell(2, 5).value, "=1+1")
        self.assertEqual(sheet.cell(2, 5).data_type, "s")
        self.assertEqual(sheet.cell(2, 11).value, "0013800000000")
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("计费重量") + 1).value, "120.00")
        selected = await self.export(selected_items=[{"source_id": "102"}, {"source_id": "106"}, {"source_id": "105"}])
        self.assertEqual(selected.max_row, 2)
        with self.assertRaises(HTTPException) as error:
            await self.export(customer_name="无匹配")
        self.assertEqual(error.exception.status_code, 400)

    async def test_export_unpaged_over_200_and_live_values_after_settlement(self):
        for source_id in range(1000, 1205):
            self.add_note(source_id)
        self.db.commit()
        sheet = await self.export(pageSize=1)
        self.assertEqual(sheet.max_row, 209)
        self.settle()
        self.db.get(ConsignmentNote, 101).form_data = json.dumps({"total_amount": "999"})
        self.db.commit()
        row = self.listing(customer_name="7")["items"][0]
        self.assertEqual(row["total_amount"], "999")
        self.assertEqual(row["agent_settlement_status"], 1)

    async def test_export_settlement_filter_matches_list_and_selected_intersection(self):
        self.settle(101)
        for status, expected_rows, expected_label in ((0, 3, "未结算"), (1, 2, "已结算")):
            sheet = await self.export(settlement_status=status, pageSize=1)
            self.assertEqual(sheet.max_row, expected_rows)
            self.assertEqual(sheet.max_row - 1, self.listing(settlement_status=status)["total"])
            self.assertTrue(all(row[1].value == expected_label for row in list(sheet.rows)[1:]))
        selected = await self.export(
            settlement_status=0, selected_items=[{"source_id": "101"}, {"source_id": "102"}],
        )
        self.assertEqual(selected.max_row, 2)
        self.assertEqual(selected.cell(2, 4).value, "普通客户")
        with self.assertRaises(HTTPException) as error:
            await self.export(settlement_status=0, selected_items=[{"source_id": "101"}])
        self.assertEqual(error.exception.status_code, 400)

    async def test_http_all_five_routes_validation_dates_and_auth(self):
        app = FastAPI()
        prefix = "/api/v1/reconciliation/road-agent"
        app.include_router(router, prefix=prefix)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_active_user] = lambda: self.user
        with TestClient(app) as client:
            schema = client.get("/openapi.json").json()
            self.assertEqual(sum(len(value) for value in schema["paths"].values()), 5)
            params = {p["name"] for p in schema["paths"][prefix+"/road"]["get"]["parameters"]}
            self.assertEqual(params, set(RoadAgentReconciliationQuery.model_fields))
            response = client.get(prefix+"/road", params={"date_start": "2026-10-19", "date_end": "2026-10-21"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["data"]["total"], 3)
            with patch("app.schemas.reconciliation_road_agent.get_china_now", return_value=datetime(2026, 10, 21, 0, 1)):
                self.assertEqual(client.get(prefix+"/road").json()["data"]["total"], 3)
            for params in ({"date_start": "bad"}, {"page": 0}, {"pageSize": 201},
                           {"settlement_status": -1}, {"settlement_status": 2}, {"settlement_status": "bad"}):
                self.assertEqual(client.get(prefix+"/road", params=params).status_code, 422)
            for value in (-1, 2, "bad"):
                self.assertEqual(client.post(prefix+"/road/export", json={"settlement_status": value}).status_code, 422)
            self.assertEqual(client.get(prefix+"/road", params={"date_start": "2026-10-22", "date_end": "2026-10-21"}).status_code, 400)
            for source_id in ("0", "bad", "9223372036854775808"):
                self.assertEqual(client.post(prefix+f"/road/{source_id}/settle").status_code, 422)
            self.assertEqual(client.post(prefix+"/road/106/settle").status_code, 404)
            self.assertEqual(client.post(prefix+"/road/101/settle").status_code, 200)
            dates = {"date_start": "2026-10-19", "date_end": "2026-10-21"}
            self.assertEqual(client.get(prefix+"/road", params={**dates, "settlement_status": 1}).json()["data"]["total"], 1)
            self.assertEqual(client.get(prefix+"/road", params={**dates, "settlement_status": 0}).json()["data"]["total"], 2)
            self.assertEqual(client.post(prefix+"/road/101/cancel-settlement").status_code, 200)
            self.assertEqual(client.post(prefix+"/road/batch-settle", json={"items": []}).status_code, 200)
            self.assertEqual(client.post(prefix+"/road/batch-settle", json={"items": [{"source_id": "-1"}]}).status_code, 422)
            self.assertEqual(client.post(prefix+"/road/export", json={
                "date_start": None, "date_end": None, "settlement_status": None,
            }).status_code, 200)
            del app.dependency_overrides[get_current_active_user]
            for method, path, body in (
                ("GET", "/road", None), ("POST", "/road/101/settle", None),
                ("POST", "/road/101/cancel-settlement", None),
                ("POST", "/road/batch-settle", {"items": []}), ("POST", "/road/export", {}),
            ):
                self.assertIn(client.request(method, prefix+path, json=body).status_code, (401, 403))


if __name__ == "__main__":
    unittest.main()
