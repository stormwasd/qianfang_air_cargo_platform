"""内部编号继续由前端传入和回显，操作记录可追溯，Excel不再导出。"""
import io
import unittest
from datetime import datetime
from types import SimpleNamespace

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.api import customer_service, cost_service
from app.api.deps import get_current_active_user
from app.database import get_db
from app.main import register_exception_handlers
from app.models.customer_service import ConsignmentInfo, ConsignmentRegistration
from app.models.cost_service import CostConsignment, CostRegistration
from app.models.consignment_operation_log import ConsignmentOperationLog


class InternalDocumentIdTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        for model in (ConsignmentInfo, ConsignmentRegistration, CostConsignment, CostRegistration, ConsignmentOperationLog):
            model.__table__.create(self.engine)
        self.db = Session(self.engine, autoflush=False)
        self.app = FastAPI()
        register_exception_handlers(self.app)
        self.app.include_router(customer_service.router, prefix="/customer")
        self.app.include_router(cost_service.router, prefix="/cost")

        async def test_db():
            return self.db

        async def test_user():
            return SimpleNamespace(id=99, name="操作员")

        self.app.dependency_overrides[get_db] = test_db
        self.app.dependency_overrides[get_current_active_user] = test_user
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        self.db.close()
        self.engine.dispose()

    async def request(self, method, path, payload=None):
        response = await self.client.request(method, path, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["data"]

    def payload(self, base, number):
        if base == "/customer":
            return {"internal_doc_id": number}
        return {"consignor_info": {"internal_doc_id": number}}

    def number_from_response(self, base, data):
        return data["internal_doc_id"] if base == "/customer" else data["consignor_info"]["internal_doc_id"]

    async def test_frontend_ids_round_trip_on_all_create_paths_and_sync(self):
        for base, source_model in (("/customer", ConsignmentInfo), ("/cost", CostConsignment)):
            for suffix in ("", "/draft"):
                number = f"FRONTEND-{base[1:]}-{suffix or 'save'}"
                payload = self.payload(base, number)
                data = await self.request("POST", base + "/consignments" + suffix, payload)
                self.assertEqual(self.number_from_response(base, data), number)
                record_id = int(data["id"])
                self.db.expire_all()
                self.assertEqual(self.db.get(source_model, record_id).internal_doc_id, number)
                for desk in ("/customer", "/cost"):
                    logs = await self.request("GET", f"{desk}/consignments/{record_id}/operation-logs")
                    self.assertEqual(logs["internal_doc_id"], number)
                    self.assertEqual(logs["consignment_id"], str(record_id))
                    self.assertEqual(logs["total"], 1)
                    self.assertNotIn("internal_doc_id", logs["items"][0])
                if suffix:
                    await self.request("PUT", f"{base}/consignments/{record_id}", payload)
                for desk, model in (("/customer", ConsignmentInfo), ("/cost", CostConsignment)):
                    self.assertEqual(self.db.get(model, record_id).internal_doc_id, number)
                    detail = await self.request("GET", f"{desk}/consignments/{record_id}")
                    self.assertEqual(self.number_from_response(desk, detail), number)
                    listing = await self.request("GET", f"{desk}/consignments")
                    item = next(item for item in listing["items"] if item["id"] == str(record_id))
                    self.assertEqual(self.number_from_response(desk, item), number)

    async def test_legacy_ids_editable_on_drafts_and_templates_as_before(self):
        self.db.add_all([
            ConsignmentInfo(id=1, status=0, internal_doc_id="YT123456789CN"),
            CostConsignment(id=2, status=0, internal_doc_id="COST-LEGACY"),
            ConsignmentRegistration(id=3, internal_doc_id="CUSTOMER-TEMPLATE"),
            CostRegistration(id=4, internal_doc_id="COST-TEMPLATE"),
        ])
        self.db.commit()
        for base, record_id in (("/customer", 1), ("/cost", 2)):
            for suffix, expected in (("/draft", "FRONTEND-DRAFT"), ("", "FRONTEND-SAVE")):
                data = await self.request("PUT", f"{base}/consignments/{record_id}{suffix}", self.payload(base, expected))
                self.assertEqual(self.number_from_response(base, data), expected)
                logs = await self.request("GET", f"{base}/consignments/{record_id}/operation-logs")
                self.assertEqual(logs["internal_doc_id"], expected)
            self.assertEqual(self.db.get(ConsignmentInfo, record_id).internal_doc_id, expected)
            self.assertEqual(self.db.get(CostConsignment, record_id).internal_doc_id, expected)
        for base, registration in (("/customer", "consignment-registration"), ("/cost", "cost-registration")):
            path = f"{base}/{registration}"
            data = await self.request("PUT", path, self.payload(base, "FRONTEND-TEMPLATE"))
            self.assertEqual(self.number_from_response(base, data), "FRONTEND-TEMPLATE")
            data = await self.request("GET", path)
            self.assertEqual(self.number_from_response(base, data), "FRONTEND-TEMPLATE")

    async def test_delete_and_batch_delete_keep_number_even_on_empty_history_page(self):
        for base in ("/customer", "/cost"):
            for batch in (False, True):
                number = f"FRONTEND-{base}-{batch}"
                data = await self.request("POST", base + "/consignments", self.payload(base, number))
                record_id = data["id"]
                await self.request("PUT", f"{base}/consignments/{record_id}/void")
                if batch:
                    await self.request("POST", base + "/consignments/batch-delete", {"ids": [record_id, record_id, "999"]})
                else:
                    await self.request("DELETE", f"{base}/consignments/{record_id}")
                self.db.expire_all()
                self.assertIsNone(self.db.get(ConsignmentInfo, int(record_id)))
                self.assertIsNone(self.db.get(CostConsignment, int(record_id)))
                for desk in ("/customer", "/cost"):
                    for page in (1, 99):
                        logs = await self.request("GET", f"{desk}/consignments/{record_id}/operation-logs?page={page}&pageSize=1")
                        self.assertEqual(logs["internal_doc_id"], number)
                        self.assertEqual(logs["total"], 3)
                        if page == 99:
                            self.assertEqual(logs["items"], [])
                snapshots = self.db.query(ConsignmentOperationLog).filter_by(consignment_id=int(record_id)).all()
                self.assertTrue(all(log.internal_doc_id == number for log in snapshots))

    async def test_old_empty_history_and_deleted_unknown_number_do_not_invent_ids(self):
        self.db.add_all([
            ConsignmentInfo(id=1, internal_doc_id="HISTORY-NUMBER"),
            CostConsignment(id=2),
            ConsignmentOperationLog(id=10, consignment_id=3, source="customer_service", action="delete",
                                    operation_name="货主委托信息删除", operator_id=99, operator_name="操作员",
                                    operated_at=datetime(2026, 10, 30)),
        ])
        self.db.commit()
        for record_id, number, total in ((1, "HISTORY-NUMBER", 0), (2, "", 0), (3, "", 1)):
            for desk in ("/customer", "/cost"):
                logs = await self.request("GET", f"{desk}/consignments/{record_id}/operation-logs")
                self.assertEqual((logs["internal_doc_id"], logs["total"]), (number, total))
        self.assertEqual((await self.client.get("/customer/consignments/999/operation-logs")).status_code, 404)

    async def test_omitted_number_is_not_generated_and_cleared_live_number_stays_empty(self):
        for base, model in (("/customer", ConsignmentInfo), ("/cost", CostConsignment)):
            for suffix in ("", "/draft"):
                data = await self.request("POST", base + "/consignments" + suffix, {})
                self.assertEqual(self.number_from_response(base, data), "")
                self.assertIsNone(self.db.get(model, int(data["id"])).internal_doc_id)
            data = await self.request("POST", base + "/consignments/draft", self.payload(base, "BEFORE-CLEAR"))
            path = f"{base}/consignments/{data['id']}"
            await self.request("PUT", path + "/draft", self.payload(base, ""))
            logs = await self.request("GET", path + "/operation-logs")
            self.assertEqual(logs["internal_doc_id"], "")

    async def test_real_exports_omit_internal_ids_without_losing_following_values(self):
        self.db.add_all([
            ConsignmentInfo(id=1, status=1, internal_doc_id="PRIVATE-CUSTOMER", customer_name="客户", pieces=8),
            CostConsignment(id=1, status=0, internal_doc_id="PRIVATE-COST", customer_name="客户", pieces=8, receivable_total=123.45),
        ])
        self.db.commit()
        for base, header_rows, count in (("/customer", 1, 132), ("/cost", 3, 137)):
            response = await self.client.post(base + "/consignments/export-excel", json={"ids": ["1"]})
            self.assertEqual(response.status_code, 200, response.text)
            workbook = load_workbook(io.BytesIO(response.content), data_only=True)
            sheet = workbook.active
            self.assertEqual(sheet.max_column, count)
            values = [cell.value for row in sheet for cell in row]
            for forbidden in ("内部单据ID", "PRIVATE-CUSTOMER", "PRIVATE-COST"):
                self.assertNotIn(forbidden, values)
            self.assertIn("客户", values)
            self.assertIn(8, values)
            if base == "/cost":
                self.assertIn(123.45, values)
            workbook.close()

    def test_openapi_preserves_input_fields_and_adds_operation_header(self):
        schemas = self.app.openapi()["components"]["schemas"]
        for name in ("ConsignmentInfoCreate", "ConsignmentInfoUpdate", "ConsignmentRegistrationSave", "ConsignorInfo"):
            self.assertIn("internal_doc_id", schemas[name]["properties"])
        self.assertIn("internal_doc_id", schemas["ConsignmentOperationLogPage"]["properties"])
