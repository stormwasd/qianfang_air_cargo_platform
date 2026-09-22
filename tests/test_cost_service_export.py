import io
import unittest
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.api.cost_service import export_cost_consignments_to_excel
from app.models.cost_service import CostConsignment
from app.schemas.cost_service import CostExportExcelRequest


class ExportQuery:
    def __init__(self, records):
        self.records = records

    def filter(self, *conditions):
        return self

    def order_by(self, *columns):
        return self

    def all(self):
        return self.records


class ExportSession:
    def __init__(self, records):
        self.records = records

    def query(self, model):
        return ExportQuery(self.records)


class CostServiceExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_subtotal_values_follow_their_group_headers_in_saved_export(self):
        record = CostConsignment(
            id=1,
            status=0,
            pay_intl_air_outsource_unit="international",
            pay_intl_air_origin="removed intl origin",
            pay_intl_air_destination="removed intl destination",
            pay_intl_air_flight_doc_no="removed intl document",
            pay_intl_air_flight_no="removed intl flight",
            pay_intl_air_flight_date=date(2026, 9, 16),
            pay_intl_air_pieces=901,
            pay_intl_air_freight_method="removed intl freight method",
            pay_intl_air_remark="removed international remark",
            pay_intl_air_subtotal=Decimal("101.25"),
            pay_trucking_outsource_unit="trucking unit",
            pay_trucking_date=date(2026, 9, 17),
            pay_trucking_pieces=902,
            pay_trucking_volume=Decimal("902.25"),
            pay_trucking_remark="removed trucking remark",
            pay_trucking_subtotal=Decimal("202.50"),
            pay_dom_air_outsource_unit="domestic unit",
            pay_dom_air_date=date(2026, 9, 18),
            pay_dom_air_origin="removed domestic origin",
            pay_dom_air_destination="removed domestic destination",
            pay_dom_air_airline="removed domestic airline",
            pay_dom_air_airline_unit="must not export",
            pay_dom_air_flight_doc_no="removed domestic document",
            pay_dom_air_flight_no="removed domestic flight",
            pay_dom_air_flight_date=date(2026, 9, 21),
            pay_dom_air_pieces=903,
            pay_dom_air_freight_method="removed domestic freight method",
            pay_dom_air_remark="removed domestic remark",
            pay_dom_air_subtotal=Decimal("303.75"),
            pay_customs_date=date(2026, 9, 19),
            pay_customs_agent="customs agent",
            pay_customs_other_fee=Decimal("904.25"),
            pay_customs_remark="customs remark",
            pay_customs_subtotal=Decimal("404.00"),
            pay_ground_date=date(2026, 9, 20),
            pay_ground_outsource_unit="ground unit",
            pay_ground_other_fee=Decimal("905.25"),
            pay_ground_remark="ground remark",
            pay_ground_subtotal=Decimal("505.25"),
            pay_total=Decimal("1516.75"),
            discount_person="discount person",
        )
        zero_record = CostConsignment(
            id=2,
            status=1,
            pay_intl_air_subtotal=0,
            pay_trucking_subtotal=0,
            pay_dom_air_subtotal=0,
            pay_customs_subtotal=0,
            pay_ground_subtotal=0,
        )
        blank_record = CostConsignment(id=3)
        response = await export_cost_consignments_to_excel(
            CostExportExcelRequest(ids=["1", "2", "3"]),
            current_user=None,
            db=ExportSession([record, zero_record, blank_record]),
        )
        workbook = load_workbook(io.BytesIO(response.body), data_only=True)
        worksheet = workbook.active
        self.addCleanup(workbook.close)

        self.assertEqual(worksheet.max_column, 90)
        self.assertEqual(worksheet.max_row, 6)
        self.assertEqual(worksheet["A1"].value, "状态")
        self.assertEqual(worksheet["A4"].value, "未提交")
        self.assertEqual(worksheet["A5"].value, "已提交")
        for subtotal_cell, expected in {
            "BB": 101.25,
            "BH": 404.00,
            "BQ": 505.25,
            "BX": 202.50,
            "CE": 303.75,
        }.items():
            with self.subTest(subtotal_cell=subtotal_cell):
                self.assertEqual(worksheet[f"{subtotal_cell}3"].value, "小计")
                self.assertEqual(worksheet[f"{subtotal_cell}4"].value, expected)
                self.assertEqual(worksheet[f"{subtotal_cell}4"].data_type, "n")
                self.assertEqual(worksheet[f"{subtotal_cell}5"].value, 0)
                self.assertIsNone(worksheet[f"{subtotal_cell}6"].value)

        expected_adjacent_values = {
            "AM4": "international",
            "BB4": 101.25,
            "BC4": "customs agent",
            "BG4": "customs remark",
            "BH4": 404.00,
            "BI4": "ground unit",
            "BP4": "ground remark",
            "BQ4": 505.25,
            "BR4": "trucking unit",
            "BX4": 202.50,
            "BY4": "domestic unit",
            "CE4": 303.75,
            "CF4": 1516.75,
            "CG4": "discount person",
        }
        for cell, expected in expected_adjacent_values.items():
            with self.subTest(cell=cell):
                self.assertEqual(worksheet[cell].value, expected)

        removed_values = {
            "removed intl origin", "removed intl destination", "removed intl document",
            "removed intl flight", "2026-09-16", 901, "removed intl freight method",
            "removed international remark", "2026-09-17", 902, 902.25,
            "removed trucking remark", "2026-09-18", "removed domestic origin",
            "removed domestic destination", "removed domestic airline", "must not export",
            "removed domestic document", "removed domestic flight", "2026-09-21", 903,
            "removed domestic freight method", "removed domestic remark", "2026-09-19",
            904.25, "2026-09-20", 905.25,
        }
        exported_values = {cell.value for cell in worksheet[4]}
        self.assertTrue(removed_values.isdisjoint(exported_values))


if __name__ == "__main__":
    unittest.main()
